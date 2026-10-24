"""Receive saved phone links, index them in WeKnora, and maintain monthly Excel ledgers."""

import argparse
from contextlib import contextmanager
import hashlib
import hmac
import ipaddress
import importlib.util
import json
import os
import secrets
import socket
import sqlite3
import threading
import time
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlsplit
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import 手机直达 as phone_bridge
import nest_paths
import weknora_client as weknora_api
from bluetooth_link import BluetoothReceiver


ROOT = Path(__file__).resolve().parent.parent / "assets"
DATA = nest_paths.ledger_dir()
DB = DATA / "captures.sqlite3"
SETTINGS = DATA / "settings.json"
PAIRING = nest_paths.pairing_file()
ENV_FILE = nest_paths.env_file()
PORT = 8793
KB_NAME = "参考与情报"
PROCESS = {'summary_enabled':False,'enable_multimodel':False,'graph_enabled':False,
           'question_generation_config':{'enabled':False}}
PHONE_CONNECTION_INTERVAL_S = 10
PHONE_CONNECTION_TIMEOUT_S = 9.5
PHONE_CONNECTION_STATE = {"state": "not_started", "started_at": 0, "finished_at": 0}
PHONE_CONNECTION_LOCK = threading.Lock()
BLUETOOTH_RECEIVER = None
HEADERS = ["收藏ID", "收藏时间", "来源", "分享文字", "链接", "入库状态", "全文状态", "WeKnora资料ID", "WeKnora网页ID", "AI处理状态", "AI反馈", "反馈时间", "备注"]
LOCK = threading.RLock()
IDENTITY_RULES = None


def capture_folder(doc):
    """File share stubs consistently without undoing a user's placement on retry."""
    current=doc.get('folder_path')or''
    defaults={'','手机收藏','手机收藏/待核实','手机收藏/待补内容'}
    if current not in defaults:
        return current
    depth=(doc.get('custom_metadata')or{}).get('content_depth')
    if doc.get('file_type')=='manual' and depth not in {'full_text','transcript','original','original_bundle'}:
        return '手机收藏/待补内容'
    return current or '手机收藏'


def annotate_capture(client, kb_id, knowledge_id, row):
    """Use the one shared identity policy; never recreate a durable capture ID."""
    global IDENTITY_RULES
    if IDENTITY_RULES is None:
        path = Path(weknora_api.__file__).resolve().parents[2] / 'ewan-weknora/scripts/material_identity.py'
        if not path.is_file():
            raise RuntimeError('收藏身份规则不可用，保留原编号待重试')
        spec = importlib.util.spec_from_file_location('capture_material_identity', path)
        IDENTITY_RULES = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(IDENTITY_RULES)
    kind = 'collection' if row['url'].startswith(('http://','https://')) else 'unclassified'
    IDENTITY_RULES.classify(client, kb_id, knowledge_id, kind)
    doc = IDENTITY_RULES.unwrap(client.request('GET','knowledge/'+quote(knowledge_id,safe='')))
    metadata = dict(doc.get('custom_metadata') or {})
    if metadata.get('phone_capture_id') not in {None,'',row['id']}:
        raise RuntimeError('收藏编号与知识原件不一致，未覆盖')
    existing_depth=metadata.get('content_depth')
    depth=existing_depth if existing_depth in {'full_text','transcript','original','original_bundle'} else ('share_text' if doc.get('file_type')=='manual' else 'unknown')
    metadata.update(phone_capture_id=row['id'],source_type='phone_capture',
                    record_role='capture' if doc.get('file_type')=='manual' else 'web_article',
                    content_depth=depth)
    IDENTITY_RULES.validate_storage(metadata)
    client.request('PUT','knowledge/'+quote(knowledge_id,safe=''),{'custom_metadata':metadata})
    destination=capture_folder(doc)
    if (doc.get('folder_path')or'')!=destination:
        client.request('POST','knowledge/folder',{'kb_id':kb_id,'knowledge_ids':[knowledge_id],'folder_path':destination})
    saved=IDENTITY_RULES.unwrap(client.request('GET','knowledge/'+quote(knowledge_id,safe='')))
    if saved.get('custom_metadata')!=metadata or saved.get('metadata')!=doc.get('metadata') or (saved.get('folder_path')or'')!=destination:
        raise RuntimeError('收藏身份回读不一致，保留原编号')


def annotate_or_defer(db, row_id, client, kb_id, knowledge_id):
    try:
        row=db.execute('SELECT * FROM captures WHERE id=?',(row_id,)).fetchone()
        annotate_capture(client, kb_id, knowledge_id, row)
        return True
    except Exception:
        db.execute("UPDATE captures SET status='分类待重试', remark='资料编号已保留；材料身份标注稍后重试' WHERE id=?", (row_id,))
        db.commit()
        return False


def annotate_page_or_defer(db, row_id, client, kb_id, page_id, url):
    """Classify recovered pages before allowing parse-only polling."""
    try:
        current = client.request('GET', 'knowledge/'+quote(page_id, safe='')).get('data', {})
        if (current.get('knowledge_base_id') != kb_id or current.get('source') != url
                or current.get('type') not in {'url', 'file_url'}):
            raise RuntimeError('网页回执与收藏链接不一致')
        # Distinct captures can share an article, but its original capture
        # attribution must remain intact. Unclaimed or same-capture pages
        # still need classification, including after a lost POST response.
        capture_id = (current.get('custom_metadata') or {}).get('phone_capture_id')
        if capture_id not in {None, '', row_id}:
            return True
        return annotate_or_defer(db, row_id, client, kb_id, page_id)
    except Exception:
        db.execute("UPDATE captures SET status='分类待重试', remark='资料编号已保留；材料身份标注稍后重试' WHERE id=?", (row_id,))
        db.commit()
        return False


def database():
    DATA.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB, timeout=20)
    db.row_factory = sqlite3.Row
    db.execute("""CREATE TABLE IF NOT EXISTS captures (
        id TEXT PRIMARY KEY, created_at INTEGER NOT NULL, source TEXT NOT NULL,
        shared_text TEXT NOT NULL, url TEXT NOT NULL, status TEXT NOT NULL,
        full_text_status TEXT NOT NULL DEFAULT '待补全文', note_id TEXT NOT NULL DEFAULT '',
        page_id TEXT NOT NULL DEFAULT '', ai_status TEXT NOT NULL DEFAULT '未处理',
        ai_feedback TEXT NOT NULL DEFAULT '', feedback_at TEXT NOT NULL DEFAULT '',
        remark TEXT NOT NULL DEFAULT '', last_attempt INTEGER NOT NULL DEFAULT 0
    )""")
    db.commit()
    return db


@contextmanager
def managed_database():
    db = database()
    try:
        with db:
            yield db
    finally:
        db.close()


def weknora_client():
    values = weknora_api.read_env_file(str(ENV_FILE)) if ENV_FILE.is_file() else {}
    base = values.get("WEKNORA_BASE_URL") or values.get("WEKNORA_HOST") or os.environ.get("WEKNORA_BASE_URL") or ""
    key = values.get("WEKNORA_API_KEY") or values.get("WEKNORA_TOKEN") or os.environ.get("WEKNORA_API_KEY") or ""
    return weknora_api.WeKnoraClient(base, key)


def payload_id(result):
    data = result.get("data", result) if isinstance(result, dict) else {}
    if not isinstance(data, dict):
        return ""
    for key in ("id", "knowledge_id", "knowledgeId"):
        if data.get(key):
            return str(data[key])
    for key in ("knowledge", "item", "document"):
        nested = data.get(key)
        if isinstance(nested, dict) and nested.get("id"):
            return str(nested["id"])
    return ""


def existing_url(client, kb_id, url):
    """Recover/reuse the same native article without merging capture events."""
    found=[]
    for number in range(1,101):
        batch=client.request('GET',f'knowledge-bases/{kb_id}/knowledge?page={number}&page_size=100').get('data',[])
        if not isinstance(batch,list):raise RuntimeError('网页列表回读异常')
        for d in batch:
            if d.get('type') in {'url','file_url'} and d.get('source')==url and d.get('knowledge_base_id')==kb_id:found.append(d)
        if len(batch)<100:break
    else:raise RuntimeError('网页列表超过核对范围')
    if len(found)>1:raise RuntimeError('相同URL有多个网页原件，暂留待核对')
    return found[0] if found else None


def capture_setting(name):
    values = weknora_api.read_env_file(str(ENV_FILE)) if ENV_FILE.is_file() else {}
    return (values.get(name) or os.environ.get(name) or "").strip()


def capture_tag_ids():
    configured = capture_setting("WEKNORA_CAPTURE_TAG_ID")
    if not configured:
        return []
    # Reject a broken source label before submitting an unlabelled capture.
    return [str(UUID(configured))]


def ensure_kb(client):
    DATA.mkdir(parents=True, exist_ok=True)
    configured = capture_setting("WEKNORA_KB_ID")
    if configured:
        kb_id = str(UUID(configured))
        result = client.request("GET", f"knowledge-bases/{quote(kb_id, safe='')}")
        base = result.get("data", {}) if isinstance(result, dict) else {}
        if base.get("id") != kb_id:
            raise RuntimeError("WeKnora 未确认配置的目标知识库")
        saved = {"kb_id": kb_id, "kb_name": base.get("name") or KB_NAME}
        temporary = SETTINGS.with_name(SETTINGS.name + "." + secrets.token_hex(6) + ".tmp")
        try:
            temporary.write_text(json.dumps(saved, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(temporary, SETTINGS)
        finally:
            temporary.unlink(missing_ok=True)
        return kb_id
    if SETTINGS.exists():
        saved = json.loads(SETTINGS.read_text(encoding="utf-8"))
        if saved.get("kb_id"):
            return saved["kb_id"]
    listing = client.request("GET", "knowledge-bases")
    bases = listing.get("data", []) if isinstance(listing, dict) else []
    if isinstance(bases, dict):
        bases = bases.get("items") or bases.get("list") or []
    for base in bases:
        if base.get("name") == KB_NAME:
            kb_id = base["id"]
            break
    else:
        raise RuntimeError('未找到已配置的参考库；保留本机收藏，不自动新建知识库')
    SETTINGS.write_text(json.dumps({"kb_id": kb_id, "kb_name": KB_NAME}, ensure_ascii=False, indent=2), encoding="utf-8")
    return kb_id


def month_path(timestamp_ms):
    month = datetime.fromtimestamp(timestamp_ms / 1000).strftime("%Y-%m")
    return DATA / f"{month}.xlsx"


def save_workbook(workbook, path):
    temporary = path.with_name(path.name + "." + secrets.token_hex(6) + ".tmp")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def ledger(timestamp_ms):
    path = month_path(timestamp_ms)
    if path.exists():
        return path
    DATA.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收藏与反馈"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:M1"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="255DD7")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [38, 20, 16, 60, 60, 24, 24, 38, 38, 18, 60, 20, 40]
    for col, width in zip("ABCDEFGHIJKLM", widths):
        sheet.column_dimensions[col].width = width
    sheet.row_dimensions[1].height = 30
    try:
        save_workbook(workbook, path)
    finally:
        workbook.close()
    return path


def sync_ledger(row):
    sync_ledgers([row])


def sync_ledgers(rows, skip_locked=False):
    months = {}
    for row in rows:
        months.setdefault(month_path(row["created_at"]), []).append(row)
    for path, monthly_rows in months.items():
        workbook = None
        try:
            ledger(monthly_rows[0]["created_at"])
            workbook = load_workbook(path)
            sheet = workbook["收藏与反馈"]
            locations = {sheet.cell(number, 1).value: number for number in range(2, sheet.max_row + 1)}
            changed = False
            for row in monthly_rows:
                target = locations.get(row["id"], sheet.max_row + 1)
                feedback = sheet.cell(target, 11).value if target <= sheet.max_row else None
                ai_status = sheet.cell(target, 10).value if target <= sheet.max_row else None
                feedback_at = sheet.cell(target, 12).value if target <= sheet.max_row else None
                user_remark = sheet.cell(target, 13).value if target <= sheet.max_row else None
                remark = str(user_remark or "")
                if row["remark"] and row["remark"] not in remark.splitlines():
                    remark = (remark + "\n" if remark else "") + row["remark"]
                values = [
                    row["id"], datetime.fromtimestamp(row["created_at"] / 1000).strftime("%Y-%m-%d %H:%M:%S"),
                    row["source"], row["shared_text"], row["url"], row["status"],
                    row["full_text_status"], row["note_id"], row["page_id"],
                    row["ai_status"] if row["ai_status"] != "未处理" else (ai_status or "未处理"),
                    row["ai_feedback"] or feedback or "", row["feedback_at"] or feedback_at or "", remark,
                ]
                if target <= sheet.max_row and all(sheet.cell(target, column).value == value
                                                   for column, value in enumerate(values, 1)):
                    continue
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(target, column, value)
                    cell.data_type = "s"
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(vertical="top", wrap_text=column in (4, 11, 13))
                locations[row["id"]] = target
                changed = True
            if changed:
                sheet.auto_filter.ref = f"A1:M{sheet.max_row}"
                save_workbook(workbook, path)
        except PermissionError:
            if not skip_locked:
                raise
        finally:
            if workbook is not None:
                workbook.close()


def receive(items):
    accepted = []
    with LOCK, managed_database() as db:
        for item in items:
            identifier = str(item.get("id", ""))
            if not identifier or len(identifier) > 80:
                continue
            text = str(item.get("shared_text", ""))[:30000]
            url = str(item.get("url", ""))[:4096]
            source = str(item.get("source", "未知应用"))[:100]
            stamp = int(item.get("created_at", 0))
            if not text or stamp < 946684800000 or stamp > int(time.time() * 1000) + 86400000:
                continue
            db.execute("INSERT OR IGNORE INTO captures(id,created_at,source,shared_text,url,status) VALUES(?,?,?,?,?,?)",
                       (identifier, stamp, source, text, url, "已到电脑，待入库"))
            accepted.append(identifier)
        db.commit()
        rows = [db.execute("SELECT * FROM captures WHERE id=?", (identifier,)).fetchone()
                for identifier in accepted]
        sync_ledgers(rows, skip_locked=True)
    return accepted


def process_one(db, row, client, kb_id):
    tag_ids = capture_tag_ids()
    identifier = row["id"]
    now = int(time.time())
    db.execute("UPDATE captures SET last_attempt=? WHERE id=?", (now, identifier))
    db.commit()
    note_id = row["note_id"]
    if not note_id:
        lines=row['shared_text'].splitlines()
        title = ((lines[0] if lines else '') or row["url"] or "手机收藏")[:90]
        content = (f"# {title}\n\n收藏时间：{datetime.fromtimestamp(row['created_at']/1000).isoformat(timespec='seconds')}"
                   f"\n来源：{row['source']}\n原始链接：{row['url'] or '无'}\n\n"
                   f"## 分享文字\n{row['shared_text']}\n\n"
                   "说明：本条记录包含手机分享文字和链接。链接页面正文是否成功抓取，请查看月度台账的「全文状态」。")
        # Recover the actual original after a lost response. A title alone is
        # never sufficient, and uncertain creation is never blindly repeated.
        candidates=[]
        for number in range(1,101):
            listing=client.request('GET',f'knowledge-bases/{kb_id}/knowledge?page={number}&page_size=100')
            batch=listing.get('data',[])
            if not isinstance(batch,list):raise RuntimeError('收藏原件列表格式异常，未创建副本')
            for d in batch:
                if d.get('file_type')=='manual' and (d.get('title')==title or d.get('custom_metadata',{}).get('phone_capture_id')==identifier):
                    actual=client.request('GET','knowledge/'+d['id']).get('data',{})
                    if actual.get('knowledge_base_id')==kb_id and actual.get('metadata',{}).get('content')==content:
                        candidates.append(actual['id'])
            if len(batch)<100:break
        else:raise RuntimeError('收藏原件列表超过核对范围，未创建副本')
        if len(candidates)>1:raise RuntimeError('发现多个相同收件原件，保留编号待核对')
        if candidates:
            note_id=candidates[0]
        else:
            if row['status']=='写入待核对':raise RuntimeError('上次写入没有明确回执，暂未找到原件；保留收件，不重复创建')
            db.execute("UPDATE captures SET status='写入待核对' WHERE id=?",(identifier,));db.commit()
            payload = {"title": title, "content": content, "status": "draft", "channel": "api",'process_config':PROCESS}
            if tag_ids:payload["tag_ids"] = tag_ids
            result = client.request("POST", f"knowledge-bases/{quote(kb_id, safe='')}/knowledge/manual", payload)
            note_id = payload_id(result)
            if not note_id:raise RuntimeError("WeKnora 未返回资料 ID；保留写入意图，不重复创建")
        db.execute("UPDATE captures SET note_id=?, status='分类待重试' WHERE id=?", (note_id, identifier))
        db.commit()
    if not annotate_or_defer(db, identifier, client, kb_id, note_id):
        return
    IDENTITY_RULES.publish_note(client,kb_id,note_id)
    # A user-abandoned web copy is never resubmitted; the link note itself stays.
    if row["url"] and not row["page_id"] and row["full_text_status"] != "网页未抓取（已放弃）":
        try:
            payload = {"url": row["url"],'process_config':PROCESS}
            if tag_ids:
                payload["tag_ids"] = tag_ids
            existing=existing_url(client,kb_id,row['url'])
            if not existing:
                if row['full_text_status']=='网页写入待核对':raise RuntimeError('网页请求结果未确认，保留回执待核对，不重复提交')
                db.execute("UPDATE captures SET full_text_status='网页写入待核对' WHERE id=?",(identifier,));db.commit()
                try:
                    result = client.request("POST", f"knowledge-bases/{quote(kb_id, safe='')}/knowledge/url", payload)
                except Exception:
                    existing=existing_url(client,kb_id,row['url'])
                    if not existing:raise
            page_id = existing['id'] if existing else payload_id(result)
            if not page_id:
                raise RuntimeError('WeKnora 未返回网页资料 ID')
            db.execute("UPDATE captures SET page_id=?, status='分类待重试', full_text_status='网页抓取已提交，待确认', remark='' WHERE id=?",
                       (page_id, identifier))
            db.commit()
            if not annotate_page_or_defer(db, identifier, client, kb_id, page_id, row['url']):
                return
        except Exception as error:
            db.execute("UPDATE captures SET full_text_status=CASE WHEN full_text_status='网页写入待核对' THEN full_text_status ELSE '待补全文' END, remark=? WHERE id=?",
                       (str(error)[:300], identifier))
    else:
        if row['page_id']:
            if not annotate_page_or_defer(db, identifier, client, kb_id, row['page_id'], row['url']):
                return
        db.execute("UPDATE captures SET full_text_status=? WHERE id=?",
                   ("无链接，保留分享文字" if not row["url"] else row["full_text_status"], identifier))
    db.commit()
    db.execute("UPDATE captures SET status='链接与分享文字已提交入库', remark=CASE WHEN remark='资料编号已保留；材料身份标注稍后重试' THEN '' ELSE remark END WHERE id=?", (identifier,))
    db.commit()
    with LOCK:
        sync_ledger(db.execute("SELECT * FROM captures WHERE id=?", (identifier,)).fetchone())


@contextmanager
def capture_write_lock():
    DATA.mkdir(parents=True,exist_ok=True)
    with (DATA/'capture-writer.lock').open('a+b') as file:
        file.seek(0)
        if not file.read(1):file.write(b'0');file.flush()
        file.seek(0)
        if os.name=='nt':
            import msvcrt
            msvcrt.locking(file.fileno(),msvcrt.LK_NBLCK,1)
        else:
            import fcntl
            fcntl.flock(file.fileno(),fcntl.LOCK_EX|fcntl.LOCK_NB)
        try:yield
        finally:
            file.seek(0)
            if os.name=='nt':msvcrt.locking(file.fileno(),msvcrt.LK_UNLCK,1)
            else:fcntl.flock(file.fileno(),fcntl.LOCK_UN)


def process_pending():
    with capture_write_lock():
        return process_pending_locked()


def process_pending_locked():
    # Remote knowledge requests must not hold the local receipt / workbook lock.
    with managed_database() as db:
        rows = db.execute("SELECT * FROM captures WHERE note_id='' OR full_text_status='网页写入待核对' OR (page_id!='' AND full_text_status='网页抓取已提交，待确认') OR status IN ('已到电脑，待入库','分类待重试','入库失败，待重试')").fetchall()
        if not rows:
            return
        client = weknora_client()
        kb_id = ensure_kb(client)
        for row in rows:
            if row["page_id"] and row["full_text_status"] == "网页抓取已提交，待确认" and row['status'] != '分类待重试':
                try:
                    result = client.request("GET", f"knowledge/{quote(row['page_id'], safe='')}")
                    page = result.get("data", result) if isinstance(result, dict) else {}
                    state = str(page.get("parse_status", "")) if isinstance(page, dict) else ""
                    if state.lower() in ("completed", "success", "processed", "published"):
                        db.execute("UPDATE captures SET full_text_status='全文已解析' WHERE id=?", (row["id"],))
                    elif state.lower() in ("failed", "error"):
                        db.execute("UPDATE captures SET full_text_status='待补全文' WHERE id=?", (row["id"],))
                    db.commit()
                    with LOCK:
                        sync_ledger(db.execute("SELECT * FROM captures WHERE id=?", (row["id"],)).fetchone())
                except Exception:
                    pass
                continue
            try:
                process_one(db, row, client, kb_id)
            except Exception as error:
                detail=str(error).replace(getattr(client,'api_key','') or '\0','[redacted]')[:300]
                db.execute("UPDATE captures SET status=CASE WHEN note_id='' AND status='写入待核对' THEN status ELSE '入库失败，待重试' END, remark=? WHERE id=?",
                           (detail, row["id"]))
                db.commit()
                try:
                    with LOCK:
                        sync_ledger(db.execute("SELECT * FROM captures WHERE id=?", (row["id"],)).fetchone())
                except PermissionError:
                    pass


def refresh_ledgers():
    with LOCK, managed_database() as db:
        sync_ledgers(db.execute("SELECT * FROM captures").fetchall(), skip_locked=True)


def maintain_phone_connection():
    global PHONE_CONNECTION_STATE
    started = time.monotonic()
    started_at = int(time.time() * 1000)
    deadline = started + PHONE_CONNECTION_TIMEOUT_S
    state = "not_connected"
    stage = "saved_address"
    address = ""
    with PHONE_CONNECTION_LOCK:
        PHONE_CONNECTION_STATE = {"state": "checking", "started_at": started_at, "finished_at": 0}
    try:
        saved = phone_bridge.config_read(deadline)
        token = saved.get("token", "")
        if not token:
            state = "unpaired"
            return
        address = saved.get("address", "")
        if address:
            try:
                reply = phone_bridge.request(address, "GET", "/heartbeat", token, timeout=3, deadline=deadline)
                if not isinstance(reply, dict) or reply.get("connected") is not True:
                    raise RuntimeError("手机未确认连接")
                state = "connected" if phone_bridge.update_pair_address(token, address, deadline) else "pairing_changed"
                return
            except RuntimeError:
                try:
                    # Keep older paired apps reachable until their update is installed.
                    reply = phone_bridge.request(address, "GET", "/status", token, timeout=3, deadline=deadline)
                    if not isinstance(reply, dict) or not reply.get("model"):
                        raise RuntimeError("手机未确认连接")
                    state = "connected" if phone_bridge.update_pair_address(token, address, deadline) else "pairing_changed"
                    return
                except Exception:
                    pass
            except Exception:
                pass
        stage = "discovery"
        found = phone_bridge.discover(token, deadline=deadline)
        if not found:
            return
        if phone_bridge.update_pair_address(token, found, deadline):
            address = found
            state = "connected"
        else:
            state = "pairing_changed"
    finally:
        elapsed = round((time.monotonic() - started) * 1000)
        if state != "connected" and time.monotonic() >= deadline:
            state = "timed_out"
        with PHONE_CONNECTION_LOCK:
            PHONE_CONNECTION_STATE = {"state": state, "stage": stage, "address": address,
                                      "started_at": started_at, "finished_at": int(time.time() * 1000),
                                      "elapsed_ms": elapsed}


def phone_connection_state():
    with PHONE_CONNECTION_LOCK:
        return dict(PHONE_CONNECTION_STATE)


def same_lan(remote):
    try:
        ip = ipaddress.ip_address(remote)
        if ip.version != 4:
            return False
        if ip.is_loopback:
            return True
        if not ip.is_private:
            return False
        return any(ip in ipaddress.IPv4Network(f"{local}/{mask}", strict=False)
                   for local, mask in phone_bridge.local_interfaces(deadline=time.monotonic() + 2))
    except (OSError, ValueError):
        return False


class Receiver(BaseHTTPRequestHandler):
    HEADER_TIMEOUT_S = 9.5
    BODY_TIMEOUT_S = 20

    def setup(self):
        super().setup()
        self.connection.settimeout(20)
        self._header_finished = threading.Event()
        self._header_expired = threading.Event()

        def expire_header():
            if not self._header_finished.is_set():
                self._header_expired.set()
                try:
                    self.connection.shutdown(socket.SHUT_RDWR)
                except OSError:
                    pass
                self.connection.close()

        self._header_watchdog = threading.Timer(self.HEADER_TIMEOUT_S, expire_header)
        self._header_watchdog.daemon = True
        self._header_watchdog.start()

    def parse_request(self):
        try:
            parsed = super().parse_request()
            return parsed and not self._header_expired.is_set()
        except OSError:
            if self._header_expired.is_set():
                return False
            raise
        finally:
            self._header_finished.set()
            self._header_watchdog.cancel()

    def finish(self):
        self._header_finished.set()
        self._header_watchdog.cancel()
        super().finish()

    def read_payload(self, length):
        def expire_body():
            try:
                self.connection.shutdown(socket.SHUT_RDWR)
            except OSError:
                pass

        deadline = threading.Timer(self.BODY_TIMEOUT_S, expire_body)
        deadline.daemon = True
        deadline.start()
        try:
            raw = self.rfile.read(length)
            if len(raw) != length:
                raise ValueError("请求内容未传完，请重试")
            return raw
        finally:
            deadline.cancel()

    def do_GET(self):
        if self.path == "/health" and self.client_address[0] in ("127.0.0.1", "::1"):
            return self.respond(200, {"service": "catdiao-nest", "version": "1.14",
                                      "phone_connection": phone_connection_state(),
                                      "bluetooth_connection": BLUETOOTH_RECEIVER.snapshot() if BLUETOOTH_RECEIVER else {"state": "not_started"}})
        if urlsplit(self.path).path == "/connection-check":
            return self.check_connection()
        if self.path not in ("/phone-bridge.apk", "/cat-diao.apk") or not same_lan(self.client_address[0]):
            return self.respond(404, {"error": "未找到"})
        apk = ROOT / "cat-diao-android-1.14.apk"
        if not apk.is_file():
            return self.respond(404, {"error": "安装包尚未生成"})
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.android.package-archive")
        self.send_header("Content-Disposition", 'attachment; filename="cat-diao.apk"')
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(apk.stat().st_size))
        self.end_headers()
        try:
            with apk.open("rb") as stream:
                while chunk := stream.read(65536):
                    self.wfile.write(chunk)
        except (BrokenPipeError, ConnectionResetError):
            pass

    def check_connection(self):
        if not same_lan(self.client_address[0]):
            return self.respond(403, {"error": "未授权"})
        try:
            token = phone_bridge.config_read().get("token", "")
            challenge = parse_qs(urlsplit(self.path).query).get("challenge", [""])[0]
            if not token or len(challenge) != 32 or any(char not in "0123456789abcdef" for char in challenge):
                return self.respond(403, {"error": "未授权"})
            expected = hmac.new(token.encode("utf-8"), ("request:" + challenge).encode("ascii"),
                                hashlib.sha256).hexdigest()
            supplied = self.headers.get("X-CatDiao-Auth", "")
            if len(supplied) != 64 or any(char not in "0123456789abcdef" for char in supplied) or not hmac.compare_digest(expected, supplied):
                return self.respond(403, {"error": "未授权"})
            proof = hmac.new(token.encode("utf-8"), ("response:" + challenge).encode("ascii"),
                             hashlib.sha256).hexdigest()
            return self.respond(200, {"service": "catdiao-nest", "connected": True},
                                {"X-CatDiao-Proof": proof})
        except (OSError, ValueError):
            return self.respond(503, {"error": "连接信息暂时不可用"})

    def do_POST(self):
        try:
            saved = phone_bridge.config_read()
            expected = saved.get("token", "")
            if self.path not in ("/captures", "/clipboard", "/file-ready") or not same_lan(self.client_address[0]) or not expected or not hmac.compare_digest(
                    self.headers.get("X-Phone-Token", ""), expected):
                return self.respond(403, {"error": "未授权"})
            length = int(self.headers.get("Content-Length", "0"))
            if length < 1 or length > (700000 if self.path == "/captures" else 262144 if self.path == "/clipboard" else 4096):
                return self.respond(413, {"error": "请求过大"})
            raw = self.read_payload(length)
            if phone_bridge.config_read().get("token") != expected:
                return self.respond(403, {"error": "配对已改变，请重新发送"})
            if self.path == "/file-ready":
                destination = phone_bridge.receive_phone_file(self.client_address[0], expected)
                return self.respond(200, {"saved": str(destination)})
            payload = json.loads(raw.decode("utf-8"))
            if self.path == "/clipboard":
                value = payload.get("text")
                if not isinstance(value, str) or not value:
                    return self.respond(400, {"error": "没有可发送的文字"})
                phone_bridge.write_windows_clipboard(value)
                return self.respond(200, {"copied": True})
            items = payload.get("items", [])
            if not isinstance(items, list) or len(items) > 20:
                return self.respond(400, {"error": "收藏格式错误"})
            accepted = receive(items)
            self.respond(200, {"accepted": accepted})
        except Exception as error:
            self.respond(502 if self.path == "/file-ready" else 500,
                         {"error": str(error)[:160]})

    def respond(self, code, payload, headers=None):
        if self.command == "POST" and self.path == "/captures":
            print(f"收藏接收请求：{self.client_address[0]}，HTTP {code}", flush=True)
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        for name, value in (headers or {}).items():
            self.send_header(name, value)
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format, *args):
        return


class ReceiverServer(ThreadingHTTPServer):
    daemon_threads = True

    def __init__(self, address, handler, max_clients=8):
        self._client_slots = threading.BoundedSemaphore(max_clients)
        super().__init__(address, handler)

    def process_request(self, request, client_address):
        if not self._client_slots.acquire(blocking=False):
            self.shutdown_request(request)
            return
        try:
            super().process_request(request, client_address)
        except BaseException:
            self._client_slots.release()
            raise

    def process_request_thread(self, request, client_address):
        try:
            super().process_request_thread(request, client_address)
        finally:
            self._client_slots.release()


def main():
    global BLUETOOTH_RECEIVER
    parser = argparse.ArgumentParser(description="手机随手收藏接收与入库")
    parser.add_argument("command", choices=["init", "listen", "sync", "status", "feedback", "giveup"])
    parser.add_argument("--id", help="收藏 ID；giveup 可用逗号分隔多个")
    parser.add_argument("--ai-status", default="已处理", help="AI 处理状态")
    parser.add_argument("--text", help="AI 的实际反馈内容；giveup 时为放弃原因")
    args = parser.parse_args()
    with managed_database() as db:
        if args.command == "status":
            print(json.dumps({"rows": db.execute("SELECT count(*) FROM captures").fetchone()[0],
                              "kb": json.loads(SETTINGS.read_text(encoding="utf-8")) if SETTINGS.exists() else None}, ensure_ascii=False))
            return
        if args.command == "feedback":
            if not args.id or not args.text:
                parser.error("feedback 需要 --id 和 --text")
            db.execute("UPDATE captures SET ai_status=?, ai_feedback=?, feedback_at=? WHERE id=?",
                       (args.ai_status, args.text, datetime.now().isoformat(timespec="seconds"), args.id))
            db.commit()
            row = db.execute("SELECT * FROM captures WHERE id=?", (args.id,)).fetchone()
            if row is None:
                parser.error("没有这条收藏 ID")
            sync_ledger(row)
            print("AI 反馈已写入月度台账")
            return
        if args.command == "giveup":
            # Stop retrying the web-page copy of captures whose link note already exists.
            ids = [item.strip() for item in (args.id or "").split(",") if item.strip()]
            if not ids:
                parser.error("giveup 需要 --id（多个用逗号分隔）")
            reason = (args.text or "网页正文多次写入失败，手动放弃重试").strip()
            remark = datetime.now().strftime("%Y-%m-%d %H:%M ") + "已放弃网页正文抓取：" + reason
            retryable = ("网页写入待核对", "网页抓取已提交，待确认", "待补全文")
            changed = []
            with capture_write_lock():
                for capture_id in ids:
                    row = db.execute("SELECT * FROM captures WHERE id=?", (capture_id,)).fetchone()
                    if row is None:
                        print(f"跳过 {capture_id}：没有这条收藏")
                        continue
                    if not row["note_id"]:
                        print(f"跳过 {capture_id}：链接资料还没入库，不能只放弃网页")
                        continue
                    if row["full_text_status"] not in retryable:
                        print(f"跳过 {capture_id}：网页状态是「{row['full_text_status']}」，无需放弃")
                        continue
                    db.execute("UPDATE captures SET full_text_status='网页未抓取（已放弃）', remark=? WHERE id=?",
                               (remark, capture_id))
                    changed.append(capture_id)
                db.commit()
                with LOCK:
                    for capture_id in changed:
                        sync_ledger(db.execute("SELECT * FROM captures WHERE id=?", (capture_id,)).fetchone())
            print(f"已停止重试 {len(changed)} 条网页正文抓取")
            return
    if args.command == "init":
        kb_id = ensure_kb(weknora_client())
        path = ledger(int(time.time() * 1000))
        print(json.dumps({"kb_id": kb_id, "ledger": str(path)}, ensure_ascii=False))
        return
    if args.command == "sync":
        process_pending()
        return
    def worker():
        while True:
            try:
                if ENV_FILE.is_file() or os.environ.get("WEKNORA_API_KEY"):
                    process_pending()
                refresh_ledgers()
            except Exception as error:
                print(f"知识库暂时无法同步：{str(error)[:160]}", flush=True)
            time.sleep(45)
    def phone_connection_worker():
        while True:
            started = time.monotonic()
            try:
                maintain_phone_connection()
            except Exception:
                pass
            time.sleep(max(0.05, PHONE_CONNECTION_INTERVAL_S - (time.monotonic() - started)))

    server = ReceiverServer(("0.0.0.0", PORT), Receiver)
    BLUETOOTH_RECEIVER = BluetoothReceiver(phone_bridge, receive)
    BLUETOOTH_RECEIVER.thread.start()
    threading.Thread(target=worker, daemon=True).start()
    threading.Thread(target=phone_connection_worker, daemon=True,
                     name="catdiao-phone-connection").start()
    print(f"手机收藏接收已开启，端口 {PORT}", flush=True)
    try:
        server.serve_forever()
    finally:
        BLUETOOTH_RECEIVER.close()
        server.server_close()


if __name__ == "__main__":
    main()
